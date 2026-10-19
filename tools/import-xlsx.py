# -*- coding: utf-8 -*-
"""Extract the word list from 单词表.xlsx into tools/xlsx-words.json.

The workbook is organized as repeated blocks:
    第1组
    (blank)
    单词 | 音标 | 词性 | 中文意思
    ...
The group heading becomes the site's List number, and numbering restarts
inside each group.

Usage:
    python tools/import-xlsx.py [path-to-xlsx]
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import openpyxl

DEFAULT_XLSX = Path(
    r"F:\xwechat_files\wxid_04j6ltb4wdnr22_1e7e\msg\file\2026-10\单词表.xlsx"
)
OUTPUT = Path(__file__).resolve().parent / "xlsx-words.json"
HEADER = ("单词", "音标", "词性", "中文意思")


def clean_phone(value: object) -> str:
    text = str(value or "").strip()
    # Excel stores /ˈeləkwənt/; the site adds the slashes itself.
    return text.strip("/").strip("[]").strip()


def main() -> int:
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not source.exists():
        print(f"找不到文件: {source}", file=sys.stderr)
        return 1

    workbook = openpyxl.load_workbook(source, data_only=True)
    sheet = workbook.active

    entries: list[dict[str, object]] = []
    current_list: int | None = None
    number = 0
    group_re = re.compile(r"^第\s*(\d+)\s*组$")

    for row_index, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=sheet.max_row, values_only=True), 1
    ):
        cells = list(row) + [None] * (4 - len(row))
        word, phone, part, meaning = cells[:4]
        filled = [value for value in cells[:4] if value not in (None, "")]

        if not filled:
            continue

        if isinstance(word, str) and all(
            cell in (None, "") for cell in cells[1:4]
        ):
            match = group_re.match(word.strip())
            if match:
                current_list = int(match.group(1))
                number = 0
                continue

        if tuple(str(cell).strip() if cell is not None else "" for cell in cells[:4]) == HEADER:
            continue

        if current_list is None:
            raise ValueError(f"第 {row_index} 行出现在任何「第N组」之前")

        word_text = str(word or "").strip()
        part_text = str(part or "").strip()
        meaning_text = str(meaning or "").strip()
        if not word_text or not part_text or not meaning_text:
            raise ValueError(f"第 {row_index} 行字段不完整: {cells[:4]!r}")

        number += 1
        phone_text = clean_phone(phone)
        entries.append(
            {
                "list": current_list,
                "number": number,
                "word": word_text,
                "part": part_text,
                "meaning": meaning_text,
                "usphone": phone_text,
                "ukphone": phone_text,
            }
        )

    if not entries:
        raise ValueError("没有读到任何词条")

    grouped: dict[int, int] = {}
    for entry in entries:
        grouped[int(entry["list"])] = grouped.get(int(entry["list"]), 0) + 1

    OUTPUT.write_text(
        json.dumps(entries, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    unique = {str(entry["word"]).lower() for entry in entries}
    print(f"源文件: {source}")
    print(f"词条: {len(entries)} 条 | 分组: {len(grouped)} 个 | 不重复单词: {len(unique)} 个")
    print("分组: " + " | ".join(f"第{k}组 {v} 条" for k, v in sorted(grouped.items())))
    print(f"写入: {OUTPUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
